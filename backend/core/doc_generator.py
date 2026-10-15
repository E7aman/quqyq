"""Генерация жалобы в .docx (python-docx)."""
import io
from datetime import date
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Cm

LABELS = {
    "ru": {"from": "От", "demand": "Прошу:", "date": "Дата", "sign": "Подпись", "sources": "Нормативные основания"},
    "kk": {"from": "Кімнен", "demand": "Сұраймын:", "date": "Күні", "sign": "Қолы", "sources": "Нормативтік негіздер"},
}


def build_docx(data, lang="ru"):
    """data: title, recipient, applicant, paragraphs[], demands[], sources[]"""
    lab = LABELS.get(lang, LABELS["ru"])
    doc = Document()
    for s in doc.sections:
        s.left_margin, s.right_margin = Cm(3), Cm(1.5)
        s.top_margin = s.bottom_margin = Cm(2)
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(13)

    head = doc.add_paragraph()
    head.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    head.paragraph_format.left_indent = Cm(7)
    head.add_run(data.get("recipient") or "[кому]").bold = True
    head.add_run(f"\n{lab['from']}: {data.get('applicant') or '[ФИО, контакты]'}")

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t.paragraph_format.space_before = Pt(18)
    run = t.add_run((data.get("title") or "Жалоба").upper())
    run.bold = True

    for text in data.get("paragraphs", []):
        p = doc.add_paragraph(text)
        p.paragraph_format.first_line_indent = Cm(1.25)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    demands = data.get("demands", [])
    if demands:
        doc.add_paragraph(lab["demand"]).runs[0].bold = True
        for i, d in enumerate(demands, 1):
            doc.add_paragraph(f"{i}. {d}")

    sources = data.get("sources", [])
    if sources:
        doc.add_paragraph().add_run(lab["sources"] + ":").italic = True
        for s in sources:
            doc.add_paragraph(f"— {s}")

    doc.add_paragraph(f"\n{lab['date']}: {date.today():%d.%m.%Y}        {lab['sign']}: ____________")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
